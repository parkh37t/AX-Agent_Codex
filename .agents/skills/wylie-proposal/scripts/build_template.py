#!/usr/bin/env python3
"""Build an editable, synthetic 12-slide Wylie-style pattern gallery.
Requires python-pptx. No source client slides, logos or font binaries are copied.
"""
import argparse
import json
from pathlib import Path
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE as S, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt
from pptx.oxml.xmlchemy import OxmlElement


def main():
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument('--out', type=Path, required=True)
    ap.add_argument('--font', default='Noto Sans CJK KR')
    args = ap.parse_args()
    if args.out.exists():
        raise SystemExit('Output exists; choose another --out path.')
    root = Path(__file__).resolve().parents[1]
    tokens = json.loads((root / 'assets/design-tokens.json').read_text())
    c = tokens['colors']
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.3333333333), Inches(7.5)

    def rect(slide, x, y, w, h, fill='white', line=None, kind=S.RECTANGLE, name='visual'):
        s = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
        s.name = name
        s.fill.solid(); s.fill.fore_color.rgb = RGBColor.from_string(c.get(fill, fill))
        if line:
            s.line.color.rgb = RGBColor.from_string(c.get(line, line)); s.line.width = Pt(1)
        else:
            s.line.fill.background()
        return s

    def text(slide, x, y, w, h, value, size=18, color='ink', bold=False, align=None, role='body'):
        s = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); s.name = role + ':' + value[:25]
        tf = s.text_frame; tf.clear(); tf.word_wrap = True
        tf.margin_left = tf.margin_right = Inches(0.02)
        tf.margin_top = tf.margin_bottom = Inches(0.02)
        for n, line in enumerate(value.split('\n')):
            p = tf.paragraphs[0] if n == 0 else tf.add_paragraph()
            p.text = line; p.font.name = args.font; p.font.size = Pt(size)
            props = p._p.get_or_add_pPr().get_or_add_defRPr()
            for script in ('ea', 'cs'):
                node = props.find('{http://schemas.openxmlformats.org/drawingml/2006/main}' + script)
                if node is None:
                    node = OxmlElement('a:' + script); props.append(node)
                node.set('typeface', args.font)
            p.font.bold = bold; p.font.color.rgb = RGBColor.from_string(c.get(color, color))
            p.space_after = Pt(8)
            if align is not None: p.alignment = align
        return s

    def arrow(slide, x, y, w=.45, h=.32):
        return rect(slide, x, y, w, h, 'purple', kind=S.RIGHT_ARROW)

    def base(kicker, title, subtitle, basis):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.background.fill.solid(); slide.background.fill.fore_color.rgb = RGBColor.from_string(c['paper'])
        text(slide, .55, .32, 11.5, .3, kicker, 11, 'purple', True, role='eyebrow')
        text(slide, .55, .72, 12.1, .65, title, 27, bold=True, role='title')
        text(slide, .55, 1.43, 12.1, .43, subtitle, 14, 'muted', role='subtitle')
        rect(slide, .55, 6.92, 12.23, .014, 'line')
        text(slide, .55, 7.04, 11.5, .22, '합성 제작 예시 · 실제 고객 성과 아님 · ' + basis, 8, 'muted', role='footer')
        text(slide, 12.1, 7.02, .65, .25, f'{len(prs.slides):02d}', 9, 'purple', align=PP_ALIGN.RIGHT, role='footer')
        slide.notes_slide.notes_text_frame.text = '이 슬라이드는 합성 템플릿입니다. 실제 고객정보·성과·가격을 포함하지 않습니다.\n참고: '+basis+'\n모든 텍스트·도형은 편집 가능합니다. 폰트는 원본과 동일하지 않은 명시적 대체 글꼴 '+args.font+'입니다.'
        return slide

    def card(slide,x,y,w,h,number,heading,body):
        rect(slide,x,y,w,h,'white',line='line')
        text(slide,x+.22,y+.20,w-.44,.35,number,14,'purple',True,role='label')
        text(slide,x+.22,y+.80,w-.44,.70,heading,21,'purple_dark',True)
        text(slide,x+.22,y+1.62,w-.44,h-1.85,body,17)

    # 1. Cover: one promise, clear scope, quiet visual.
    s=base('WYLIE / PROPOSAL PATTERNS','고객의 문제를 실행 가능한 제안으로 연결합니다','개인 제안서 제작 스킬 · 편집 가능한 레이아웃 12종','D24:p1,6 / 합성 재구성')
    rect(s,0,0,13.3333,6.92,'purple_dark',name='cover-background')
    text(s,.75,.65,9,.35,'WYLIE  ·  PROPOSAL STUDIO',14,'white',True,role='eyebrow')
    text(s,.75,1.70,10.8,1.7,'고객의 문제를\n실행 가능한 제안으로',36,'white',True,role='title')
    text(s,.78,4.0,9.6,.9,'진단 → 전략 → 경험 → 실행 → 검증\n한 줄 메시지를 근거와 산출물로 연결합니다.',20,'white')
    text(s,.78,5.72,10,.5,'합성 예시  |  폰트·색·내용을 고객 규격에 맞춰 조정',14,'white')

    # 2. Executive strategy.
    s=base('01 / 핵심 요약','세 가지 실행으로 고객의 반복 확인을 줄입니다','문제와 해법, 검증할 결과를 같은 순서로 연결합니다.','D24:p3,17-18 / D03:p8')
    for i,(h,b) in enumerate([('기준을 고정','가격·권한·상태 기준을\n거래 이력과 연결합니다.'),('흐름을 연결','입력부터 승인까지\n예외 경로를 포함합니다.'),('성과를 확인','처리시간·재작업률을\n동일 기준으로 비교합니다.')]):
        card(s,.55+i*4.15,2.18,3.93,3.95,f'0{i+1}',h,b)

    # 3. Diagnosis.
    s=base('02 / 진단','현재의 불편을 원인과 검증 질문으로 분해합니다','관찰 사실과 가설을 구분해 해결 방향을 정합니다.','D26:p13,15 / 합성 재구성')
    for i,(h,b) in enumerate([('관찰한 현상','같은 항목을 반복 확인\n[자료로 확인할 내용]'),('원인 가설','기준과 이력이 분리됨\n[인터뷰로 검증할 내용]'),('검증할 변화','한 화면에서 근거 확인\n[시나리오로 시험할 내용]')]):
        card(s,.55+i*4.15,2.18,3.93,3.95,f'0{i+1}',h,b)

    # 4. Strategy house: hierarchy rather than decoration.
    s=base('03 / 전략 체계','하나의 목표를 세 전략과 공통 기반으로 구체화합니다','목표 → 전략 → 과제 → 산출물의 수직 추적을 유지합니다.','D24:p19 / D12:p65 / 합성 재구성')
    rect(s,1.1,2.12,11.15,.65,'purple_dark',kind=S.CHEVRON)
    text(s,1.45,2.24,10.45,.4,'목표: 근거가 보이는 빠른 업무 처리',20,'white',True,PP_ALIGN.CENTER)
    for i,(h,b) in enumerate([('정확성','정책·데이터 기준\n검증 규칙 목록'),('연결성','역할별 업무 여정\n핵심 화면 시안'),('실행력','단계별 전환 계획\n검수·이관 증거')]):
        rect(s,1.1+i*3.85,3.02,3.45,2.2,'white',line='line')
        text(s,1.32+i*3.85,3.24,3,.5,h,23,'purple',True,PP_ALIGN.CENTER)
        text(s,1.32+i*3.85,4.0,3,1,b,17,align=PP_ALIGN.CENTER)
    rect(s,1.1,5.52,11.15,.7,'purple_light')
    text(s,1.3,5.66,10.75,.4,'공통 기반  ·  근거 관리 / 디자인시스템 / 책임과 검수',17,'purple_dark',True,PP_ALIGN.CENTER)

    # 5. Before/after on the same comparison axes.
    s=base('04 / 전환 논리','동일한 비교축으로 변화의 의미를 설명합니다','우수해 보이는 말 대신 실제 바뀌는 처리 방식을 씁니다.','D24:p24 / D26:p17')
    for x,h in [(.55,'현재'),(6.9,'제안')]:
        rect(s,x,2.15,5.87,.65,'purple_dark' if x>1 else 'muted')
        text(s,x+.2,2.28,5.47,.35,h,20,'white',True)
    for j,(a,b) in enumerate([('입력: 여러 화면을 왕복','입력: 역할별 시작점 통합'),('판단: 담당자에게 기준 확인','판단: 기준·이력을 화면에 표시'),('확인: 결과만 사후 점검','확인: 예외와 증거를 함께 기록')]):
        y=3.05+j*1.04
        rect(s,.55,y,5.87,.83,'white');text(s,.78,y+.20,5.41,.43,a,17)
        arrow(s,6.5,y+.23,.28,.25)
        rect(s,6.9,y,5.87,.83,'purple_light');text(s,7.13,y+.20,5.41,.43,b,17)

    # 6. Journey.
    s=base('05 / 경험 설계','고객 여정의 마찰을 화면과 측정 지점으로 연결합니다','단계 수보다 단계별 역할·불편·개입·증거가 중요합니다.','D26:p13 / D11:시각 사례')
    for i,(h,b) in enumerate([('탐색','필요를 발견\n검색 성공률'),('확인','근거를 비교\n정보 확인시간'),('실행','예외를 처리\n과업 완료율'),('유지','결과를 재사용\n재작업률')]):
        x=.55+i*3.12
        rect(s,x,2.6,2.88,.65,'purple');text(s,x+.1,2.73,2.68,.4,h,21,'white',True,PP_ALIGN.CENTER)
        text(s,x+.2,3.65,2.48,1.2,b,18,align=PP_ALIGN.CENTER)
        if i<3:arrow(s,x+2.94,2.82,.16,.20)
    rect(s,.55,5.65,12.23,.62,'purple_light');text(s,.8,5.81,11.73,.35,'아래에 화면 ID·수용조건·계측 이벤트를 연결합니다.',17,'purple_dark',True)

    # 7. Evidence matrix.
    s=base('06 / 평가 대응','모든 약속은 산출물과 검수 근거를 가집니다','조견표를 목차·화면·일정과 함께 갱신합니다.','D03:p10 / D25:p44-46')
    xs=[.55,2.63,6.20,9.66]; widths=[2.08,3.57,3.46,3.12]
    headers=['요구 / 평가','대응 전략','제출 산출물','검수 근거']
    rows=[['REQ-예시-01','입력 예외의 일관 처리','화면·규칙 정의','예외 시나리오'],['REQ-예시-02','기준과 변경 이력 연결','데이터 매핑표','정합성 검증'],['REQ-예시-03','단계별 운영 이관','전환·교육 계획','완료 체크리스트']]
    for j,h in enumerate(headers):rect(s,xs[j],2.22,widths[j],.68,'purple_dark');text(s,xs[j]+.14,2.4,widths[j]-.28,.4,h,17,'white',True)
    for i,row in enumerate(rows):
        for j,v in enumerate(row):
            rect(s,xs[j],3.02+i*.94,widths[j],.84,'white' if i%2==0 else 'purple_light',line='line')
            text(s,xs[j]+.14,3.24+i*.94,widths[j]-.28,.48,v,15)

    # 8. Architecture with named boundaries.
    s=base('07 / 구조','고객 채널과 업무 규칙, 데이터 책임을 분리합니다','배포 구성도가 아니라 책임 경계를 보여 주는 개념도입니다.','D14:p7,31-36 / 합성 재구성')
    for j,(label,blocks) in enumerate([('채널',['고객 앱','업무 포털','관리 화면']),('업무 서비스',['정책·권한','과업·승인','알림·조회']),('데이터 / 연계',['기준정보','거래 이력','외부 시스템'])]):
        y=2.16+j*1.36;text(s,.65,y+.35,1.55,.5,label,15,'purple_dark',True)
        for i,b in enumerate(blocks):
            rect(s,2.3+i*3.48,y,3.18,.93,'white' if j!=1 else 'purple_light',line='line')
            text(s,2.48+i*3.48,y+.26,2.82,.47,b,19,align=PP_ALIGN.CENTER)
    text(s,2.4,6.36,10.1,.3,'연계 방향·호출 주체·오류 책임은 상세 설계에서 명시합니다.',12,'muted',role='label')

    # 9. Delivery roadmap, phases are not arbitrary percentages.
    s=base('08 / 실행 계획','단계마다 승인할 산출물과 다음 진입 조건을 정합니다','기간은 합성 예시이며 실제 범위·인력·의존조건으로 다시 산정합니다.','D24:p13-14 / D29:단계별 추진')
    stages=[('정의','요구·정책 합의','기준 승인'),('설계','화면·데이터 계약','시나리오 승인'),('구축','기능·통합 시험','수용 기준 통과'),('전환','교육·운영 이관','운영 준비 확인')]
    for i,(h,a,b) in enumerate(stages):
        x=.55+i*3.12;rect(s,x,2.5,2.87,.66,'purple')
        text(s,x+.15,2.65,2.57,.4,h,20,'white',True)
        text(s,x+.16,3.48,2.55,.85,a,18)
        rect(s,x,4.72,2.87,.90,'purple_light');text(s,x+.17,4.94,2.53,.52,b,16,'purple_dark',True)
        if i<3:arrow(s,x+2.92,2.74,.2,.2)

    # 10. UI proof without copying client screens.
    s=base('09 / UX 증명','핵심 화면 한 장으로 전략이 작동하는 방식을 보여줍니다','시안은 구현·테스트 완료와 구분하고 검증할 과업을 붙입니다.','D11:이미지형 UI 사례 / D26:핵심 UX')
    rect(s,.55,2.16,7.48,4.33,'white',line='line');rect(s,.55,2.16,7.48,.57,'purple_dark')
    text(s,.76,2.28,7.06,.35,'합성 업무 화면  ·  목업',16,'white',True)
    rect(s,.80,3.02,6.98,.7,'paper');text(s,1.0,3.19,6.58,.4,'조건을 입력하면 기준과 결과를 함께 확인',16)
    for j in range(3):
        rect(s,.8,4.0+j*.62,6.98,.47,'purple_light' if j==0 else 'paper')
    text(s,8.43,2.35,4.05,.7,'설계 의도',22,'purple_dark',True)
    text(s,8.43,3.30,4.05,2.1,'① 반복 확인을 줄이는 정보\n② 현재 상태와 다음 행동\n③ 예외 처리와 변경 이력',18)
    text(s,8.43,5.73,4.05,.6,'검증: 담당자가 과업을 완료하는가',14,'purple',True)

    # 11. Credentials as a capability bridge.
    s=base('10 / 왜 와일리인가','과거 경험을 이번 과제에서 맡을 책임으로 바꿉니다','로고 나열 대신 유사 조건·기여 범위·전이 방법·근거를 제시합니다.','D02:레퍼런스 / D24:p17 / D03:p9')
    for i,(h,b) in enumerate([('유사한 상황','여러 역할이 함께 처리하는\n복잡한 업무 경험'),('우리의 기여','고객 여정과 업무 규칙을\n화면·검수로 연결'),('이번의 적용','대표 과업을 먼저 검증하고\n전환·운영까지 인계')]):
        card(s,.55+i*4.15,2.18,3.93,3.95,f'0{i+1}',h,b)

    # 12. Decision close.
    s=base('11 / 의사결정','착수 전에 세 가지 기준을 함께 확정합니다','좋은 제안은 마지막에 다음 행동을 명확히 남깁니다.','D24:p48-50 / 합성 재구성')
    for i,(h,b) in enumerate([('범위','이번 단계의 필수 과업과\n후속 과제를 나눕니다.'),('기준','정책·데이터·검수의\n정본과 책임자를 정합니다.'),('일정','검토·승인·전환의\n의존조건을 확인합니다.')]):
        card(s,.55+i*4.15,2.18,3.93,3.95,f'0{i+1}',h,b)

    args.out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(args.out)
    print(json.dumps({'output':str(args.out),'slides':len(prs.slides),'font':args.font,'synthetic':True},ensure_ascii=False))


if __name__=='__main__': main()
