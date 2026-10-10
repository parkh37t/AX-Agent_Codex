#!/usr/bin/env python3
"""Conservative static PPTX audit. Does not certify rendering or factual truth.
Requires python-pptx. Exit 0=static checks pass, 2=fail, 3=tool error.
"""
import argparse
import csv
import json
import re
import sys
from pathlib import Path
from pptx import Presentation

CLAIM_STATES = {'verified_fact', 'source_claim', 'proposal', 'hypothesis', 'unresolved', 'placeholder'}
PLACEHOLDER = re.compile(r'삽입\s*예정|작성\s*필요|미회신|\[고객명\]|\[사업명\]|<[^<>\n]{1,80}>|\b(?:TODO|TBD|TBC|Lorem ipsum)\b', re.I)


def audit(path, mode, profile, plan=None, evidence=None):
    prs = Presentation(path)
    issues, warnings, notes = [], [], []
    explicit_fonts, unresolved_font_runs = set(), 0
    body_min = 16 if profile == 'presentation' else 10
    tolerance = 9144  # 0.01in; rounding only, not arbitrary overflow tolerance.
    if not prs.slides:
        issues.append({'check':'nonempty', 'message':'No slides.'})
    for number, slide in enumerate(prs.slides, 1):
        slide_text = []
        for shape in slide.shapes:
            # Top-level rectangles are in slide coordinates. Group children are
            # not tested as global coordinates without transform composition.
            if shape.left < -tolerance or shape.top < -tolerance or shape.left+shape.width > prs.slide_width+tolerance or shape.top+shape.height > prs.slide_height+tolerance:
                issues.append({'slide':number,'check':'bounds','shape':shape.name,'message':'Top-level shape outside slide; inspect cropping/intent.'})
            stack = [shape]
            while stack:
                node = stack.pop()
                if hasattr(node, 'shapes'):
                    stack.extend(node.shapes)
                frames = []
                if node.has_text_frame: frames.append(node.text_frame)
                if node.has_table:
                    frames.extend(c.text_frame for row in node.table.rows for c in row.cells)
                for frame in frames:
                    if frame.text.strip(): slide_text.append(frame.text)
                    role = node.name.split(':')[0]
                    if role == 'footer': minimum=7
                    elif role in ['label','eyebrow']: minimum=10
                    else: minimum=body_min
                    for paragraph in frame.paragraphs:
                        for run in paragraph.runs:
                            if not run.text.strip(): continue
                            face = run.font.name or paragraph.font.name
                            font_size = run.font.size if run.font.size is not None else paragraph.font.size
                            if face: explicit_fonts.add(face)
                            else: unresolved_font_runs+=1
                            if font_size is not None and font_size.pt < minimum:
                                warnings.append({'slide':number,'check':'small_text','shape':node.name,'size_pt':round(font_size.pt,2),'minimum_pt':minimum})
        text='\n'.join(slide_text)
        found=sorted(set(PLACEHOLDER.findall(text)))
        if found:
            entry={'slide':number,'check':'placeholder','values':found}
            (issues if mode=='final' else warnings).append(entry)
        if not text.strip(): warnings.append({'slide':number,'check':'text_layer_absent','message':'Image-only or empty; inspect visually.'})
        notes_frame = slide.notes_slide.notes_text_frame if slide.has_notes_slide else None
        if notes_frame is not None:
            found=sorted(set(PLACEHOLDER.findall(notes_frame.text)))
            if found: warnings.append({'slide':number,'check':'speaker_note_placeholder','values':found})
    if plan:
        with open(plan,encoding='utf-8-sig',newline='') as f: rows=list(csv.DictReader(f))
        ids=[r.get('slide_id','').strip() for r in rows]
        if not rows or any(not x for x in ids) or len(ids)!=len(set(ids)):
            issues.append({'check':'slide_plan_ids','message':'Slide IDs must be present and unique.'})
        for r in rows:
            sid=r.get('slide_id','')
            if not r.get('governing_message','').strip(): issues.append({'check':'slide_plan_message','slide_id':sid})
            state=r.get('claim_status','')
            if state not in CLAIM_STATES: issues.append({'check':'claim_status','slide_id':sid,'value':state})
            if mode=='final' and state in ['unresolved','placeholder']: issues.append({'check':'unresolved_final_claim','slide_id':sid})
            if state in ['verified_fact','source_claim'] and not r.get('evidence_ids','').strip(): issues.append({'check':'missing_evidence_reference','slide_id':sid})
        if len(rows)!=len(prs.slides): issues.append({'check':'slide_plan_count','plan_rows':len(rows),'slides':len(prs.slides)})
    if evidence:
        with open(evidence,encoding='utf-8-sig',newline='') as f: rows=list(csv.DictReader(f))
        ids=[r.get('evidence_id','').strip() for r in rows]
        if not rows or any(not x for x in ids) or len(ids)!=len(set(ids)):
            issues.append({'check':'evidence_ids','message':'Evidence IDs must be present and unique.'})
        for r in rows:
            state=r.get('claim_status','')
            if state not in CLAIM_STATES: issues.append({'check':'evidence_claim_status','evidence_id':r.get('evidence_id')})
            if state=='verified_fact' and not r.get('verification_method','').strip(): issues.append({'check':'verified_without_method','evidence_id':r.get('evidence_id')})
    notes.extend(['Run and paragraph-default sizes/names are checked; layout/master/theme inheritance is not resolved.',
                  'No automatic certification of text overflow, overlap, colors, facts, arithmetic, RFP completeness, chart semantics or font availability.',
                  'Group child geometry requires manual inspection after transform; cropping can be intentional.',
                  'Warnings require human/agent review even when static status is pass.'])
    return {'status':'fail' if issues else 'pass','scope':'static_checks_only','slides':len(prs.slides),'profile':profile,'mode':mode,
            'failures':issues,'warnings':warnings,'explicit_font_families':sorted(explicit_fonts),
            'runs_without_explicit_latin_font':unresolved_font_runs,'limitations':notes,
            'rendering':'not_run','factual_verification':'not_run'}


def main():
    ap=argparse.ArgumentParser(description=__doc__)
    ap.add_argument('input',type=Path)
    ap.add_argument('--mode',choices=['draft','final'],default='draft')
    ap.add_argument('--profile',choices=['report','presentation'],default='report')
    ap.add_argument('--slide-plan',type=Path)
    ap.add_argument('--evidence',type=Path)
    ap.add_argument('--json-out',type=Path)
    args=ap.parse_args()
    if args.json_out and args.json_out.exists(): raise SystemExit('Output exists; choose another --json-out.')
    try:
        result=audit(args.input,args.mode,args.profile,args.slide_plan,args.evidence)
    except Exception as exc:
        result={'status':'error','error_type':type(exc).__name__,'message':str(exc)}
    payload=json.dumps(result,ensure_ascii=False,indent=2)
    if args.json_out:
        args.json_out.parent.mkdir(parents=True,exist_ok=True);args.json_out.write_text(payload,encoding='utf-8')
    print(payload)
    return 0 if result['status']=='pass' else 2 if result['status']=='fail' else 3


if __name__=='__main__': sys.exit(main())
