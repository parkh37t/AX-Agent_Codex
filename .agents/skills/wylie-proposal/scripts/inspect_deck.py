#!/usr/bin/env python3
"""Inspect authorized PPTX/PDF/DOCX files without executing embedded content.

Requires python-pptx, PyMuPDF, python-docx. Font declarations are NOT equivalent
to effective rendered fonts; inherited/default faces are explicitly unresolved.
"""
import argparse
import collections
import hashlib
import json
import re
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

NS = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
EMU = 914400


def xml_text(data):
    root = ET.fromstring(data)
    return [n.text for n in root.iter() if n.tag.endswith("}t") and n.text]


def inspect_pptx(path):
    from pptx import Presentation
    prs = Presentation(path)
    fonts, sizes, colors = collections.Counter(), collections.Counter(), collections.Counter()
    pages = []
    inherited_runs = 0

    def runs_from_shape(shape):
        if shape.has_text_frame:
            yield from (r for p in shape.text_frame.paragraphs for r in p.runs)
        if shape.has_table:
            for row in shape.table.rows:
                for cell in row.cells:
                    yield from (r for p in cell.text_frame.paragraphs for r in p.runs)

    def shape_info(shape, parent=None):
        nonlocal inherited_runs
        texts = []
        if shape.has_text_frame:
            texts.extend(p.text for p in shape.text_frame.paragraphs if p.text.strip())
        if shape.has_table:
            texts.extend(" | ".join(c.text for c in row.cells) for row in shape.table.rows)
        for run in runs_from_shape(shape):
            weight = max(1, len(run.text.strip()))
            declared = []
            for tag in ("latin", "ea", "cs"):
                for node in run._r.findall("a:rPr/a:" + tag, NS):
                    face = node.get("typeface")
                    if face:
                        declared.append(face)
            for face in set(declared):
                fonts[face] += weight
            if not declared:
                inherited_runs += 1
            if run.font.size is not None:
                sizes[str(round(run.font.size.pt, 2))] += weight
            for tag in ("srgbClr", "schemeClr"):
                for node in run._r.findall(".//a:" + tag, NS):
                    colors[tag + ":" + node.get("val", "")] += weight
        item = {
            "name": shape.name, "shape_type": str(shape.shape_type),
            "x_in": round(shape.left / EMU, 4), "y_in": round(shape.top / EMU, 4),
            "w_in": round(shape.width / EMU, 4), "h_in": round(shape.height / EMU, 4),
            "coordinate_scope": "slide" if parent is None else "group-local:" + parent,
            "text": "\n".join(texts),
        }
        if shape.has_chart:
            chart = shape.chart
            item["chart_type"] = str(chart.chart_type)
            item["chart_series"] = [s.name for s in chart.series]
        try:
            image = shape.image
        except AttributeError:
            image = None
        except ValueError:
            image = None
            item["image_status"] = "linked_or_unavailable_embedded_image"
        if image is not None:
            item["image"] = {"sha1": image.sha1, "size_px": image.size}
        if hasattr(shape, "shapes"):
            item["children"] = [shape_info(c, shape.name) for c in shape.shapes]
        return item

    def collect_text(shapes):
        for shape in shapes:
            if shape["text"]:
                yield shape["text"]
            if shape.get("chart_series"):
                yield "[chart series] " + " | ".join(shape["chart_series"])
            yield from collect_text(shape.get("children", []))

    for number, slide in enumerate(prs.slides, 1):
        shapes = [shape_info(s) for s in slide.shapes]
        notes_frame = slide.notes_slide.notes_text_frame if slide.has_notes_slide else None
        notes = notes_frame.text if notes_frame is not None else ""
        pages.append({"page": number, "text": "\n".join(collect_text(shapes)),
                      "notes": notes, "shapes": shapes})
    themes, diagrams, embedded = {}, {}, []
    with zipfile.ZipFile(path) as z:
        for name in z.namelist():
            if name.startswith("ppt/theme/") and name.endswith(".xml"):
                root = ET.fromstring(z.read(name))
                faces = sorted({n.get("typeface") for n in root.iter() if n.get("typeface")})
                palette = {}
                for scheme in root.findall(".//a:clrScheme", NS):
                    for node in scheme:
                        child = next(iter(node), None)
                        if child is not None:
                            palette[node.tag.split("}")[-1]] = child.get("val") or child.get("lastClr")
                themes[name] = {"declared_faces_including_unused_scripts": faces, "palette": palette}
            if name.startswith("ppt/diagrams/") and name.endswith(".xml"):
                text = xml_text(z.read(name))
                if text:
                    diagrams[name] = text
            if name.startswith("ppt/fonts/"):
                embedded.append(name)
    return {"type": "pptx", "width_in": prs.slide_width / EMU,
            "height_in": prs.slide_height / EMU, "pages": pages,
            "explicit_font_character_counts": dict(fonts.most_common()),
            "explicit_size_pt_character_counts": dict(sizes.most_common()),
            "explicit_text_color_character_counts": dict(colors.most_common()),
            "runs_without_explicit_typeface": inherited_runs,
            "font_scope": "Run-level OOXML declarations weighted by run characters, deduplicated within a run. Different script faces may share a run. Inheritance/effective rendering unresolved.",
            "themes": themes, "diagram_text": diagrams, "embedded_font_parts": embedded}


def inspect_pdf(path):
    import fitz
    doc = fitz.open(path)
    pages, fonts, sizes, colors = [], collections.Counter(), collections.Counter(), collections.Counter()
    for i, page in enumerate(doc, 1):
        spans = []
        for block in page.get_text("dict")["blocks"]:
            for line in block.get("lines", []):
                for span in line.get("spans", []):
                    weight = max(1, len(span["text"].strip()))
                    fonts[span["font"]] += weight
                    sizes[str(round(span["size"], 2))] += weight
                    colors[f"#{span['color']:06X}"] += weight
                    spans.append({k: span[k] for k in ["text", "font", "size", "bbox", "color"]})
        pages.append({"page": i, "text": page.get_text(sort=True), "notes": "",
                      "width_pt": page.rect.width, "height_pt": page.rect.height,
                      "spans": spans, "images": len(page.get_images()),
                      "font_resources": [list(x) for x in page.get_fonts()]})
    return {"type": "pdf", "pages": pages, "pdf_font_character_counts": dict(fonts.most_common()),
            "pdf_size_pt_character_counts": dict(sizes.most_common()),
            "pdf_color_character_counts": dict(colors.most_common()),
            "font_scope": "Original PDF text-span font names; image-contained text requires visual inspection."}


def inspect_docx(path):
    from docx import Document
    doc = Document(path)
    text = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        text.extend(" | ".join(c.text for c in row.cells) for row in table.rows)
    with zipfile.ZipFile(path) as z:
        supplementary = {}
        font_declarations = {}
        for name in z.namelist():
            if re.match(r"word/(?:header|footer|footnotes|endnotes|comments).*\.xml$", name):
                supplementary[name] = xml_text(z.read(name))
            if name.startswith('word/') and name.endswith('.xml'):
                tree = ET.fromstring(z.read(name))
                faces = set()
                for node in tree.iter():
                    if node.tag.endswith('}rFonts'):
                        for key, value in node.attrib.items():
                            if key.split('}')[-1] in ('ascii', 'hAnsi', 'eastAsia', 'cs'):
                                faces.add(value)
                    if name=='word/fontTable.xml' and node.tag.endswith('}font'):
                        faces.update(v for k,v in node.attrib.items() if k.endswith('}name'))
                if faces: font_declarations[name] = sorted(faces)
    return {"type": "docx", "pagination": "not_resolved", "supplementary": supplementary,
            "font_scope": "DOCX part-level font declarations including defaults/styles/fontTable; not effective rendered text or proof of actual use.",
            "font_declarations": font_declarations,
            "pages": [{"page": 1, "text": "\n".join(text), "notes": ""}]}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("input", type=Path)
    parser.add_argument("--out", required=True, type=Path, help="Dedicated output folder; existing outputs are never overwritten")
    args = parser.parse_args()
    path, out = args.input, args.out
    signature = path.read_bytes()[:8]
    if signature[:4] in [bytes.fromhex("00051607"), bytes.fromhex("00051600")]:
        raise SystemExit("AppleDouble/AppleSingle metadata, not a presentation. Supply the real document.")
    if any((out / name).exists() for name in ("analysis.json", "content.md")):
        raise SystemExit("Output already exists; choose a fresh --out directory.")
    ext = path.suffix.lower()
    func = {".pptx": inspect_pptx, ".pdf": inspect_pdf, ".docx": inspect_docx}.get(ext)
    if not func:
        raise SystemExit("Supported inputs: .pptx, .pdf, .docx")
    data = func(path)
    data.update({"filename": path.name, "sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
                 "page_count": None if ext=='.docx' else len(data["pages"]),
                 "content_unit_count": len(data["pages"]), "extracted_not_visually_verified": True})
    out.mkdir(parents=True, exist_ok=True)
    (out / "analysis.json").write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
    lines = ["# " + path.name, "", "Extraction is not a claim of visual or factual verification."]
    for page in data["pages"]:
        lines.extend(["", "## " + ("Document body" if ext == ".docx" else f"Page {page['page']}"), "", page["text"]])
        if page["notes"].strip():
            lines.extend(["", "### Speaker notes", page["notes"]])
    for name, texts in data.get("diagram_text", {}).items():
        lines.extend(["", "## Diagram XML: " + name, "\n".join(texts)])
    for name, texts in data.get("supplementary", {}).items():
        lines.extend(["", "## Supplementary: " + name, "\n".join(texts)])
    (out / "content.md").write_text("\n".join(lines), encoding="utf-8")
    print(json.dumps({"file": path.name, "pages": data["page_count"], "output": str(out)}, ensure_ascii=False))


if __name__ == "__main__":
    main()
